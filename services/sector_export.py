"""Export Excel du TCD secteur (format mi-parcours CRIC EXT — Secteur CIME)."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def _rate_color(rate):
    if rate >= 70:
        return '166534'
    if rate >= 40:
        return '1e40af'
    if rate >= 20:
        return '92400e'
    return '991b1b'


def build_sector_tcd_workbook(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = 'TCD Secteur CIME'

    headers = [
        'SECTEUR', 'DECLARANT', 'Taux de declaration', 'DEFAIL',
        'Taux de DEFAIL', 'NEANT', 'Taille du fichier',
    ]
    header_fill = PatternFill('solid', fgColor='1e40af')
    header_font = Font(bold=True, color='FFFFFF', size=10)
    thin = Side(style='thin', color='93c5fd')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border

    for r_idx, row in enumerate(rows, 2):
        values = [
            row.get('SECTEUR', ''),
            row.get('DECLARANT', 0),
            row.get('Taux de declaration', 0),
            row.get('DEFAIL', 0),
            row.get('Taux de DEFAIL', 0),
            row.get('NEANT', 0),
            row.get('Taille du fichier', 0),
        ]
        is_total = str(row.get('SECTEUR', '')).lower().startswith('total')
        taux_decl = float(row.get('Taux de declaration', 0) or 0)
        color = _rate_color(taux_decl) if not is_total else '111827'

        for c_idx, val in enumerate(values, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.border = border
            cell.alignment = Alignment(horizontal='left' if c_idx == 1 else 'right', vertical='center')
            if c_idx in (3, 5):
                cell.number_format = '0.0"%"'
            if is_total:
                cell.font = Font(bold=True, color=color)
            elif c_idx == 1:
                cell.font = Font(color=color, bold=taux_decl >= 70)
            elif c_idx == 3:
                cell.font = Font(color=color, bold=True)

        if is_total:
            for c_idx in range(1, len(headers) + 1):
                ws.cell(row=r_idx, column=c_idx).fill = PatternFill('solid', fgColor='eff6ff')

    widths = [28, 14, 18, 12, 16, 10, 16]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    return wb


def sector_tcd_to_bytes(rows):
    bio = BytesIO()
    build_sector_tcd_workbook(rows).save(bio)
    bio.seek(0)
    return bio


def build_sous_secteur_tcd_workbook(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = 'TCD Sous secteur'
    headers = [
        'SOUS SECTEUR', 'Taille', 'declaration', 'Taux de declaration',
        'def', 'NEANT', 'Non contributeurs', 'Taux de non contributeurs',
    ]
    header_fill = PatternFill('solid', fgColor='0f766e')
    header_font = Font(bold=True, color='FFFFFF', size=10)
    thin = Side(style='thin', color='cbd5e1')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border

    for r_idx, row in enumerate(rows, 2):
        values = [
            row.get('SOUS SECTEUR', ''),
            row.get('Taille', 0),
            row.get('declaration', 0),
            row.get('Taux de declaration', 0),
            row.get('def', 0),
            row.get('NEANT', 0),
            row.get('Non contributeurs', 0),
            row.get('Taux de non contributeurs', 0),
        ]
        is_total = str(row.get('SOUS SECTEUR', '')).lower().startswith('total')
        for c_idx, val in enumerate(values, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.border = border
            cell.alignment = Alignment(
                horizontal='left' if c_idx == 1 else 'right',
                vertical='center',
            )
            if c_idx in (4, 8):
                cell.number_format = '0.0"%"'
            if is_total:
                cell.font = Font(bold=True)
        if is_total:
            for c_idx in range(1, len(headers) + 1):
                ws.cell(row=r_idx, column=c_idx).fill = PatternFill('solid', fgColor='ecfdf5')

    widths = [32, 12, 14, 20, 10, 10, 18, 24]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    return wb


def sous_secteur_tcd_to_bytes(rows):
    bio = BytesIO()
    build_sous_secteur_tcd_workbook(rows).save(bio)
    bio.seek(0)
    return bio


def build_ca_sous_secteur_workbook(rows, col_current='CA 2026', col_previous='CA 2025'):
    wb = Workbook()
    ws = wb.active
    ws.title = 'CA par sous-secteur'
    headers = ['SOUS SECTEUR', col_current, col_previous, 'Ecart absolu', 'Evolution']
    header_fill = PatternFill('solid', fgColor='c2410c')
    header_font = Font(bold=True, color='FFFFFF', size=10)
    thin = Side(style='thin', color='fdba74')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border

    for r_idx, row in enumerate(rows, 2):
        is_total = row.get('_is_total') or str(row.get('SOUS SECTEUR', '')).upper() == 'TOTAL'
        evo = row.get('Evolution')
        evo_display = f'{evo}%' if evo is not None else '—'
        values = [
            row.get('SOUS SECTEUR', ''),
            row.get(col_current, 0),
            row.get(col_previous, 0),
            row.get('Ecart absolu', 0),
            evo_display if not isinstance(evo, (int, float)) else evo,
        ]
        for c_idx, val in enumerate(values, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.border = border
            cell.alignment = Alignment(horizontal='left' if c_idx == 1 else 'right', vertical='center')
            if c_idx in (2, 3, 4):
                cell.number_format = '#,##0'
            if c_idx == 5 and isinstance(evo, (int, float)):
                cell.number_format = '0.0"%"'
            if is_total:
                cell.font = Font(bold=True)
        if is_total:
            for c_idx in range(1, len(headers) + 1):
                ws.cell(row=r_idx, column=c_idx).fill = PatternFill('solid', fgColor='fff7ed')

    widths = [36, 16, 16, 16, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    return wb


def ca_sous_secteur_to_bytes(rows, col_current='CA 2026', col_previous='CA 2025'):
    bio = BytesIO()
    build_ca_sous_secteur_workbook(rows, col_current, col_previous).save(bio)
    bio.seek(0)
    return bio


def build_reliquataires_sous_secteur_workbook(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Reliquataires SS'
    headers = [
        'SOUS SECTEUR', 'Reliquataires', 'Montant reliquataires',
        'Emissions', 'Ratio des reliquataires',
    ]
    header_fill = PatternFill('solid', fgColor='ea580c')
    header_font = Font(bold=True, color='FFFFFF', size=10)
    thin = Side(style='thin', color='fdba74')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border

    for r_idx, row in enumerate(rows, 2):
        is_total = row.get('_is_total') or str(row.get('SOUS SECTEUR', '')).lower().startswith('total')
        values = [
            row.get('SOUS SECTEUR', ''),
            row.get('Reliquataires', 0),
            row.get('Montant reliquataires', 0),
            row.get('Emissions', 0),
            row.get('Ratio des reliquataires', 0),
        ]
        for c_idx, val in enumerate(values, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.border = border
            cell.alignment = Alignment(horizontal='left' if c_idx == 1 else 'right', vertical='center')
            if c_idx in (3, 4):
                cell.number_format = '#,##0'
            if c_idx == 5:
                cell.number_format = '0.0"%"'
            if is_total:
                cell.font = Font(bold=True)
        if is_total:
            for c_idx in range(1, len(headers) + 1):
                ws.cell(row=r_idx, column=c_idx).fill = PatternFill('solid', fgColor='fff7ed')

    widths = [36, 14, 20, 18, 22]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    return wb


def reliquataires_sous_secteur_to_bytes(rows):
    bio = BytesIO()
    build_reliquataires_sous_secteur_workbook(rows).save(bio)
    bio.seek(0)
    return bio
