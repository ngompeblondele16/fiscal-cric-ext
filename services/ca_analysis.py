"""Analyse CA par sous-secteur — import vectorisé, RECHERCHEV, agrégation N / N-1."""
import json
import os
import unicodedata
from datetime import datetime

import pandas as pd

from models import Contribuable, CaAnalysisMeta, CaBySousSecteur, CaNiuLine
from import_utils import (
    _patch_openpyxl_european_numbers,
    _normalize_col,
    _standardize_columns,
    _header_has_niu,
    _dataframe_from_raw_with_header,
    _score_import_sheet,
    STANDARD_COLUMNS,
)

# Les 3 colonnes du fichier CA additionnées pour obtenir le Total CA.
# Total CA = CA_TAXABLE + CA_EXPORT + CA_EXONERE
CA_AMOUNT_COLUMNS = ('ca_taxable', 'ca_export', 'ca_exonere')
# Ordre d'affichage imposé des sous-secteurs (les libellés inconnus passent en bas, avant le total).
_CA_SOUS_SECTEUR_ORDER = [
    'Agriculture',
    'Sylviculture',
    'Industrie Extractive',
    'Industries Agroalimentaires',
    'Industries Manufasturières',
    'Eau et Electricité',
    'Construction',
    'Commerce Général',
    'Prestation de Service',
    'Vente de Boisson',
    'Restaurant et Hébergement',
    'Transport et communication',
    'Banques et Assurance',
    'Services marchands',
    'Administration Publique',
    'Education',
    'Service non marchands',
]
_CA_SOUS_SECTEUR_ORDER_INDEX = {
    ''.join(
        c for c in unicodedata.normalize('NFKD', label.lower().strip())
        if not unicodedata.combining(c)
    ): idx
    for idx, label in enumerate(_CA_SOUS_SECTEUR_ORDER)
}
_CA_ORDER_ALIASES = {
    'industries manufacturieres': 'industries manufasturieres',
    'services non marchands': 'service non marchands',
    'service marchands': 'services marchands',
}


def has_ca_columns(cols):
    # Vrai si les 3 colonnes CA sont présentes dans le fichier.
    return all(c in (cols or set()) for c in CA_AMOUNT_COLUMNS)


def clear_ca_analysis(db):
    # Vide toutes les tables liées à l'analyse CA.
    db.query(CaNiuLine).delete()
    db.query(CaBySousSecteur).delete()
    db.query(CaAnalysisMeta).delete()


def ca_analysis_ready(db):
    # Vrai si un tableau CA a déjà été calculé.
    return db.query(CaBySousSecteur).count() > 0


def _parse_year(val, default=None):
    # Extrait une année (int) d'une valeur libre ('2026', '2026-01', etc.).
    if val is None:
        return default
    s = str(val).strip()
    if not s:
        return default
    try:
        y = int(float(s))
        if 1990 <= y <= 2100:
            return y
    except (ValueError, TypeError):
        pass
    for part in s.replace('/', '-').split('-'):
        part = part.strip()
        if len(part) == 4 and part.isdigit():
            return int(part)
    return default


def _vectorize_amounts(series):
    # Convertit une colonne texte en montants numériques (gère espaces et virgule décimale).
    s = series.astype(str).str.strip()
    s = s.str.replace('\xa0', '', regex=False).str.replace(' ', '', regex=False)
    s = s.str.replace(',', '.', regex=False)
    return pd.to_numeric(s, errors='coerce').fillna(0.0)


def _normalize_niu_series(series):
    # Uniformise les NIU de la colonne (majuscules, sans espaces) pour fiabiliser la jointure.
    return (
        series.astype(str)
        .str.strip()
        .str.upper()
        .str.replace(' ', '', regex=False)
    )


def _normalize_niu_value(value):
    # Même normalisation NIU, sur une valeur unique.
    if value is None:
        return ''
    return str(value).strip().upper().replace(' ', '')


def _normalize_order_key(value):
    # Clé de tri/fusion d'un sous-secteur (minuscules, sans accents, espaces réduits, alias).
    raw = ' '.join((value or '').strip().lower().split())
    norm = ''.join(
        c for c in unicodedata.normalize('NFKD', raw)
        if not unicodedata.combining(c)
    )
    # Règle robuste : toutes les variantes « (services) marchands » fusionnent,
    # quelle que soit la casse, les espaces ou le singulier/pluriel.
    if 'marchand' in norm:
        return 'service non marchands' if 'non marchand' in norm else 'services marchands'
    return _CA_ORDER_ALIASES.get(norm, norm)


# clé normalisée -> libellé officiel d'affichage
_CA_ORDER_CANONICAL = {
    _normalize_order_key(label): label
    for label in _CA_SOUS_SECTEUR_ORDER
}


def _canonical_sous_secteur(label):
    """
    Fusionne les variantes d'un même sous-secteur (casse, accents, pluriel).
    Ex.: « service marchands » / « Services marchands » -> « Services marchands ».
    Un libellé vide devient « #N/A ».
    """
    clean = (label or '').strip()
    if not clean:
        return '#N/A'
    return _CA_ORDER_CANONICAL.get(_normalize_order_key(clean), clean)


def _sort_sous_secteurs_for_ca(labels):
    # Trie selon l'ordre imposé ; les libellés hors liste finissent en bas.
    return sorted(
        labels,
        key=lambda s: (
            _CA_SOUS_SECTEUR_ORDER_INDEX.get(_normalize_order_key(s), 10_000),
            _normalize_order_key(s),
        ),
    )


def _niu_column_candidates():
    # Noms de colonnes acceptés comme NIU (pour ne lire que les colonnes utiles).
    aliases = {'niu'}
    aliases.update(STANDARD_COLUMNS.get('niu', []))
    return aliases


def _read_ca_dataframe_fast(path):
    """
    Lit uniquement les colonnes utiles du fichier CA (NIU + 3 montants),
    pour réduire fortement I/O et mémoire sur gros volumes.
    """
    needed = set(CA_AMOUNT_COLUMNS)
    niu_aliases = _niu_column_candidates()

    def _usecol(col):
        nc = _normalize_col(str(col))
        return (nc in needed) or (nc in niu_aliases)

    ext = os.path.splitext(path)[1].lower()
    if ext == '.csv':
        try:
            df = pd.read_csv(path, dtype=str, encoding='utf-8-sig', usecols=_usecol)
        except UnicodeDecodeError:
            df = pd.read_csv(path, dtype=str, encoding='latin-1', usecols=_usecol)
    else:
        xl = pd.ExcelFile(path)
        try:
            preview, sheet, _score = _load_best_excel_sheet_preview(xl)
            if preview is None or preview.empty:
                return pd.DataFrame()
            df = pd.read_excel(xl, sheet_name=sheet, dtype=str, usecols=_usecol)
        finally:
            xl.close()

    if df is None or df.empty:
        return pd.DataFrame()

    norm_map = {orig: _normalize_col(orig) for orig in df.columns}
    df = df.rename(columns=norm_map)
    df = _standardize_columns(df)
    required = {'niu', *CA_AMOUNT_COLUMNS}
    if not required.issubset(set(df.columns)):
        return pd.DataFrame()
    return df[['niu', *CA_AMOUNT_COLUMNS]].copy()


def _load_ca_series_vectorized(filepaths):
    """
    Lecture vectorisée — évite iterrows(), ~10–50× plus rapide sur gros fichiers.
    Retourne une Series index=NIU, values=Total CA.

    Process Excel reproduit :
      1) colonne Total CA = CA_TAXABLE + CA_EXPORT + CA_EXONERE (par ligne),
      2) TCD du fichier CA : NIU en lignes, Total CA en valeurs (somme),
      3) cette table sert de source à la RECHERCHEV sur le fichier en ligne.
    """
    frames = []
    for path in filepaths or []:
        if not path or not os.path.isfile(path):
            continue
        df = _read_ca_dataframe_fast(path)
        if df is None or df.empty:
            continue
        work = df.copy()
        work['niu'] = _normalize_niu_series(work['niu'])
        for col in CA_AMOUNT_COLUMNS:
            work[col] = _vectorize_amounts(work[col])
        work['total_ca'] = work[list(CA_AMOUNT_COLUMNS)].sum(axis=1)
        # Aligner RECHERCHEV Excel: pas de filtre de forme NIU trop strict.
        # On ne garde que les clés non vides pour la jointure.
        work = work[work['niu'].notna()]
        work = work[~work['niu'].isin(['', 'NAN', 'NONE', 'NULL'])]
        if not work.empty:
            frames.append(work[['niu', 'total_ca']])

    if not frames:
        return pd.Series(dtype=float)

    combined = pd.concat(frames, ignore_index=True)
    # TCD Excel : une ligne par NIU, Total CA = somme des lignes du fichier CA.
    tcd = combined.groupby('niu', as_index=False)['total_ca'].sum()
    return tcd.set_index('niu')['total_ca']


def _contrib_lookup(db):
    """NIU → secteur / sous-secteur en une requête."""
    secteur = {}
    sous_secteur = {}
    for niu, sec, ss in db.query(Contribuable.niu, Contribuable.secteur, Contribuable.sous_secteur).all():
        niu_key = _normalize_niu_value(niu)
        if not niu_key:
            continue
        if sec and str(sec).strip():
            secteur[niu_key] = str(sec).strip()
        ss_clean = (ss or '').strip()
        if ss_clean and ss_clean.upper() not in ('#N/A', 'N/A', 'NA'):
            sous_secteur[niu_key] = ss_clean
    return secteur, sous_secteur


def _ca_statut(niu, in_ca, sous_secteur):
    # Étiquette de contrôle : NIU absent du fichier CA, sans sous-secteur, ou OK.
    if not in_ca:
        return '#N/A CA'
    if not sous_secteur:
        return '#N/A sous-secteur'
    return 'OK'


def _build_niu_lines(db, ca_cur, ca_prev, secteur_map, sous_map):
    """
    Base = fichier en ligne (une ligne par NIU en ligne) + RECHERCHEV CA.

    Reproduit exactement le process Excel :
      - TCD préalable sur chaque fichier CA (NIU → somme Total CA),
      - on se place dans le fichier en ligne,
      - colonne CA N   = RECHERCHEV(NIU ; TCD CA N   ; Total CA ; FAUX),
      - colonne CA N-1 = RECHERCHEV(NIU ; TCD CA N-1 ; Total CA ; FAUX),
      - #N/A si le NIU n'est pas trouvé dans le TCD CA.
    Retourne la liste des lignes (dict) prêtes pour l'insertion.
    """
    ca_cur_dict = ca_cur.to_dict() if ca_cur is not None and len(ca_cur) else {}
    ca_prev_dict = ca_prev.to_dict() if ca_prev is not None and len(ca_prev) else {}

    en_ligne_rows = db.query(
        Contribuable.niu,
        Contribuable.raison,
        Contribuable.montant_declare,
    ).all()

    lines = []
    for niu, raison, montant_declare in sorted(en_ligne_rows, key=lambda r: (r[0] or '')):
        niu_key = _normalize_niu_value(niu)
        in_cur = niu_key in ca_cur_dict
        in_prev = niu_key in ca_prev_dict
        raw_ss = sous_map.get(niu_key, '')
        ss = _canonical_sous_secteur(raw_ss)  # fusionne les variantes ('' -> '#N/A')
        lines.append({
            'niu': niu,
            'raison': raison or '',
            'montant_declare': float(montant_declare or 0),
            'secteur': secteur_map.get(niu_key, ''),
            'sous_secteur': ss,
            'ca_current': float(ca_cur_dict.get(niu_key) or 0) if in_cur else None,
            'ca_previous': float(ca_prev_dict.get(niu_key) or 0) if in_prev else None,
            'statut_ca': _ca_statut(niu_key, in_cur, raw_ss),
        })
    return lines


def _aggregate_from_niu_lines(niu_lines):
    """
    TCD depuis le fichier en ligne : somme du Total CA par sous-secteur.

    - on part des lignes du fichier en ligne (RECHERCHEV déjà appliquée),
    - on retire les #N/A CA (NIU absent du fichier CA → ca None, non sommé),
    - un NIU en ligne sans sous-secteur alimente la ligne « #N/A ».
    Retourne (agg_current, agg_previous, mapped_current, excluded_current).
    """
    agg_cur = {}
    agg_prev = {}
    mapped = 0
    excluded = 0
    for ln in niu_lines:
        ss = (ln.get('sous_secteur') or '').strip() or '#N/A'
        cur = ln.get('ca_current')
        prev = ln.get('ca_previous')
        if cur is not None:
            agg_cur[ss] = agg_cur.get(ss, 0.0) + float(cur or 0)
            mapped += 1
        else:
            excluded += 1
        if prev is not None:
            agg_prev[ss] = agg_prev.get(ss, 0.0) + float(prev or 0)
    return agg_cur, agg_prev, mapped, excluded


def run_ca_analysis(db, current_paths, previous_paths=None, year_current=None, year_previous=None, user='admin'):
    # Pipeline complet : lecture fichiers CA -> RECHERCHEV sur le fichier en ligne
    # -> TCD par sous-secteur -> écart/évolution N vs N-1, le tout stocké en base.
    secteur_map, sous_map = _contrib_lookup(db)
    if not secteur_map and not sous_map:
        return {'success': False, 'error': 'Fichier en ligne mis a jour introuvable — terminez les etapes 4 et 5.'}
    if not sous_map:
        return {
            'success': False,
            'error': 'Aucun sous-secteur disponible — terminez l\'étape Secteurs (étape 5).',
        }

    ca_current = _load_ca_series_vectorized(current_paths)
    if len(ca_current) == 0:
        return {'success': False, 'error': 'Aucune ligne CA valide dans le(s) fichier(s) importé(s).'}

    ca_previous = _load_ca_series_vectorized(previous_paths or [])

    y_cur = year_current or datetime.utcnow().year
    y_prev = year_previous if year_previous is not None else (y_cur - 1)

    # Base = fichier en ligne + RECHERCHEV CA (process Excel), puis TCD depuis ces lignes.
    clear_ca_analysis(db)
    niu_lines = _build_niu_lines(db, ca_current, ca_previous, secteur_map, sous_map)
    if niu_lines:
        db.bulk_insert_mappings(CaNiuLine, niu_lines)

    agg_cur, agg_prev, mapped_cur, excl_cur = _aggregate_from_niu_lines(niu_lines)

    if not agg_cur and not agg_prev:
        return {
            'success': False,
            'error': 'Aucun NIU du fichier en ligne trouvé dans le(s) fichier(s) CA (tous #N/A).',
        }

    # Une ligne par sous-secteur + une ligne total, avec écart absolu et évolution %.
    all_ss = _sort_sous_secteurs_for_ca(set(agg_cur) | set(agg_prev))
    rows = []
    total_cur = total_prev = 0.0

    for ss in all_ss:
        cur = float(agg_cur.get(ss, 0) or 0)
        prev = float(agg_prev.get(ss, 0) or 0)
        ecart = cur - prev
        if prev:
            evolution = round((ecart / prev) * 100, 1)
        elif cur:
            evolution = None
        else:
            evolution = 0.0
        rows.append(CaBySousSecteur(
            sous_secteur=ss,
            ca_current=cur,
            ca_previous=prev,
            ecart_absolu=ecart,
            evolution_pct=evolution,
            is_total=0,
        ))
        total_cur += cur
        total_prev += prev

    total_ecart = total_cur - total_prev
    total_evo = round((total_ecart / total_prev) * 100, 1) if total_prev else None
    rows.append(CaBySousSecteur(
        sous_secteur='TOTAL',
        ca_current=total_cur,
        ca_previous=total_prev,
        ecart_absolu=total_ecart,
        evolution_pct=total_evo,
        is_total=1,
    ))

    meta = CaAnalysisMeta(
        year_current=y_cur,
        year_previous=y_prev,
        mapped_count=mapped_cur,
        excluded_na_count=excl_cur,
        lines_current=len(ca_current),
        lines_previous=len(ca_previous),
        source_json=json.dumps({
            'current_files': [os.path.basename(p) for p in (current_paths or []) if p],
            'previous_files': [os.path.basename(p) for p in (previous_paths or []) if p],
        }, ensure_ascii=False),
        user=user,
    )
    db.add(meta)
    for r in rows:
        db.add(r)

    return {
        'success': True,
        'year_current': y_cur,
        'year_previous': y_prev,
        'mapped': mapped_cur,
        'excluded': excl_cur,
        'sous_secteurs': len(all_ss),
        'total_ca_current': total_cur,
        'total_ca_previous': total_prev,
    }


def get_ca_meta(db):
    # Dernière analyse CA enregistrée (années, compteurs, fichiers source).
    return db.query(CaAnalysisMeta).order_by(CaAnalysisMeta.id.desc()).first()


def ca_table_rows(db):
    # Construit le tableau CA affiché (fusion des variantes, tri, total).
    meta = get_ca_meta(db)
    if not meta:
        return [], None

    y_cur = meta.year_current or datetime.utcnow().year
    y_prev = meta.year_previous or (y_cur - 1)
    col_cur = f'CA {y_cur}'
    col_prev = f'CA {y_prev}'

    db_rows = db.query(CaBySousSecteur).all()
    total_rows = [r for r in db_rows if bool(r.is_total)]

    # Fusion des variantes de libellé (casse/pluriel/accents) sur une seule ligne.
    merged = {}
    for r in db_rows:
        if bool(r.is_total):
            continue
        label = _canonical_sous_secteur(r.sous_secteur)
        acc = merged.get(label)
        if acc is None:
            merged[label] = {'cur': float(r.ca_current or 0), 'prev': float(r.ca_previous or 0)}
        else:
            acc['cur'] += float(r.ca_current or 0)
            acc['prev'] += float(r.ca_previous or 0)

    ordered_labels = sorted(
        merged.keys(),
        key=lambda s: (
            _CA_SOUS_SECTEUR_ORDER_INDEX.get(_normalize_order_key(s), 10_000),
            _normalize_order_key(s),
        ),
    )

    rows = []
    for label in ordered_labels:
        cur = merged[label]['cur']
        prev = merged[label]['prev']
        ecart = cur - prev
        if prev:
            evo = round((ecart / prev) * 100, 1)
        elif cur:
            evo = None
        else:
            evo = 0.0
        rows.append({
            'SOUS SECTEUR': label,
            col_cur: cur,
            col_prev: prev,
            'Ecart absolu': ecart,
            'Evolution': evo,
            '_is_total': False,
        })

    for r in total_rows:
        rows.append({
            'SOUS SECTEUR': r.sous_secteur,
            col_cur: r.ca_current,
            col_prev: r.ca_previous,
            'Ecart absolu': r.ecart_absolu,
            'Evolution': r.evolution_pct,
            '_is_total': True,
        })
    return rows, {
        'year_current': y_cur,
        'year_previous': y_prev,
        'col_current': col_cur,
        'col_previous': col_prev,
        'meta': meta,
    }


def get_ca_top10_dashboard(db):
    """Top 10 CA sous-secteur (plus élevés et plus faibles) pour le dashboard."""
    rows, info = ca_table_rows(db)
    if not rows or not info:
        return None

    col_cur = info['col_current']
    year = info['year_current']
    normal_rows = [r for r in rows if not r.get('_is_total')]
    if not normal_rows:
        return None

    normalized = []
    for r in normal_rows:
        val = float(r.get(col_cur) or 0)
        normalized.append({
            'label': r.get('SOUS SECTEUR') or '#N/A',
            'value': val,
        })

    top_high = sorted(normalized, key=lambda x: x['value'], reverse=True)[:10]
    top_low = sorted(normalized, key=lambda x: x['value'])[:10]
    max_high = max((x['value'] for x in top_high), default=0.0)
    max_low = max((x['value'] for x in top_low), default=0.0)

    def _with_pct(items, max_val):
        out = []
        for item in items:
            pct = round((item['value'] / max_val) * 100, 1) if max_val > 0 else 0
            out.append({
                'label': item['label'],
                'value': item['value'],
                'pct': pct,
            })
        return out

    return {
        'year': year,
        'top_high': _with_pct(top_high, max_high),
        'top_low': _with_pct(top_low, max_low),
    }


def ca_negative_gap_details(db, sous_secteur):
    """Détail NIU d'un sous-secteur pour expliquer un écart absolu négatif.

    Reprend tel quel le fichier en ligne (une ligne par NIU du sous-secteur)
    après ajout des colonnes CA 2026 / CA 2025 (RECHERCHEV).
    - tous les NIU du sous-secteur sont listés, même sans CA (#N/A affiché « — ») ;
    - écart de ligne = CA 2026 − CA 2025, le #N/A comptant comme 0 ;
    - les totaux ne somment que les CA réellement présents (= tableau agrégé).
    """
    meta = get_ca_meta(db)
    if not meta or not sous_secteur:
        return [], None

    y_cur = meta.year_current or datetime.utcnow().year
    y_prev = meta.year_previous or (y_cur - 1)
    col_cur = f'CA {y_cur}'
    col_prev = f'CA {y_prev}'

    target = _canonical_sous_secteur(sous_secteur)
    lines = [
        ln for ln in db.query(CaNiuLine).all()
        if _canonical_sous_secteur(ln.sous_secteur) == target
    ]
    rows = []
    sum_cur = 0.0
    sum_prev = 0.0
    for ln in lines:
        # Toutes les lignes du sous-secteur sont conservées (fichier en ligne complet).
        cur_raw = ln.ca_current
        prev_raw = ln.ca_previous
        cur_num = float(cur_raw) if cur_raw is not None else None
        prev_num = float(prev_raw) if prev_raw is not None else None
        # Écart : année absente traitée comme 0 pour expliquer le TCD (comme agrégation séparée N / N-1)
        ecart = (cur_num or 0.0) - (prev_num or 0.0)
        if cur_num is not None:
            sum_cur += cur_num
        if prev_num is not None:
            sum_prev += prev_num
        rows.append({
            'NIU': ln.niu,
            'Raison': ln.raison or '',
            col_cur: cur_num,
            col_prev: prev_num,
            'Ecart absolu': ecart,
            'statut_ca': ln.statut_ca or '',
        })

    rows = sorted(rows, key=lambda r: (r['Ecart absolu'], r['NIU'] or ''))
    summary = {
        'sous_secteur': sous_secteur,
        'col_current': col_cur,
        'col_previous': col_prev,
        'total_current': sum_cur,
        'total_previous': sum_prev,
        'total_ecart': sum_cur - sum_prev,
        'nb_lignes': len(rows),
    }
    return rows, summary


def consolide_ca_export_rows(db):
    """Fichier en ligne mis a jour + RECHERCHEV CA + sous-secteur (verification avant aggregation)."""
    meta = get_ca_meta(db)
    if not meta:
        return [], None

    y_cur = meta.year_current or datetime.utcnow().year
    y_prev = meta.year_previous or (y_cur - 1)
    col_cur = f'Total CA {y_cur}'
    col_prev = f'Total CA {y_prev}'

    lines = db.query(CaNiuLine).order_by(CaNiuLine.niu).all()
    rows = []
    for ln in lines:
        rows.append({
            'NIU': ln.niu,
            'Nom/Raison sociale': ln.raison or '',
            'Montant à payer': round(float(ln.montant_declare or 0), 2),
            'Secteur': ln.secteur or '',
            'Sous secteur': ln.sous_secteur or '#N/A',
            col_cur: ln.ca_current if ln.ca_current is not None else '#N/A',
            col_prev: ln.ca_previous if ln.ca_previous is not None else '#N/A',
            'Statut CA': ln.statut_ca or '',
        })
    return rows, {'col_current': col_cur, 'col_previous': col_prev, 'meta': meta}


def inspect_ca_upload_fast(filepath):
    """
    Validation upload CA sans charger tout le classeur en mémoire (aperçu en-têtes + comptage rapide).
    Les autres étapes utilisent inspect_upload_file (lecture complète) ; les fichiers CA sont souvent volumineux.
    """
    _patch_openpyxl_european_numbers()
    ext = os.path.splitext(filepath)[1].lower()
    errors = []

    if ext == '.csv':
        try:
            preview = pd.read_csv(filepath, dtype=str, nrows=200, encoding='utf-8-sig')
        except UnicodeDecodeError:
            preview = pd.read_csv(filepath, dtype=str, nrows=200, encoding='latin-1')
        norm_map = {orig: _normalize_col(orig) for orig in preview.columns}
        preview = preview.rename(columns=norm_map)
        preview = _standardize_columns(preview)
        cols = set(preview.columns)
        try:
            with open(filepath, 'rb') as f:
                nb_rows = max(0, sum(1 for _ in f) - 1)
        except OSError:
            nb_rows = len(preview)
        ftype = 'ca' if has_ca_columns(cols) else 'unknown'
    else:
        xl = pd.ExcelFile(filepath)
        df, _sheet, _score = _load_best_excel_sheet_preview(xl)
        if df is None or df.empty:
            cols = set()
            nb_rows = 0
            ftype = 'unknown'
        else:
            norm_map = {orig: _normalize_col(orig) for orig in df.columns}
            df = df.rename(columns=norm_map)
            df = _standardize_columns(df)
            cols = set(df.columns)
            ftype = 'ca' if has_ca_columns(cols) else detect_ca_or_unknown(cols)
            nb_rows = _count_excel_rows_fast(filepath, _sheet)
            if nb_rows <= 0 and _sheet:
                # Fallback robuste (ex: .xls / moteurs non supportés par openpyxl).
                # On lit une seule colonne pour éviter une lecture complète coûteuse.
                try:
                    fallback_df = pd.read_excel(filepath, sheet_name=_sheet, dtype=str, usecols=[0])
                    nb_rows = max(0, len(fallback_df))
                except Exception:
                    pass
        xl.close()

    if 'niu' not in cols:
        errors.append('Colonne NIU obligatoire pour le fichier CA.')
    if not has_ca_columns(cols):
        missing = [c.upper() for c in CA_AMOUNT_COLUMNS if c not in cols]
        errors.append('Colonnes CA obligatoires manquantes : ' + ', '.join(missing))

    return {
        'file_type': ftype,
        'nb_rows': nb_rows,
        'errors': errors,
        'ok': len(errors) == 0,
    }


def detect_ca_or_unknown(cols):
    # Type de fichier d'après les colonnes présentes.
    if has_ca_columns(cols):
        return 'ca'
    return 'unknown'


def _load_best_excel_sheet_preview(xl, preview_rows=120):
    """Repère la bonne feuille via aperçu uniquement (pas de chargement intégral)."""
    best_sheet = None
    best_score = -1
    best_header_row = 0
    best_use_raw = False
    best_df = pd.DataFrame()

    for sheet in xl.sheet_names:
        try:
            df_std = pd.read_excel(xl, sheet_name=sheet, dtype=str, nrows=preview_rows)
            if not df_std.empty:
                score = _score_import_sheet(df_std)
                if score > best_score:
                    best_score = score
                    best_sheet = sheet
                    best_df = df_std
                    best_header_row = 0
                    best_use_raw = False
        except Exception:
            pass
        try:
            raw = pd.read_excel(xl, sheet_name=sheet, dtype=str, header=None, nrows=preview_rows + 15)
            if raw.empty:
                continue
            for hr in range(min(15, len(raw))):
                if not _header_has_niu(raw.iloc[hr].tolist()):
                    continue
                df_hdr = _dataframe_from_raw_with_header(raw, hr)
                score = _score_import_sheet(df_hdr)
                if score > best_score:
                    best_score = score
                    best_sheet = sheet
                    best_df = df_hdr
                    best_header_row = hr
                    best_use_raw = True
        except Exception:
            pass

    return best_df, best_sheet, best_score


def _count_excel_rows_fast(filepath, sheet_name=None):
    """Compte les lignes via openpyxl read_only (beaucoup plus rapide que pandas sur gros xlsx)."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        try:
            if sheet_name and sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
            else:
                ws = wb.active
            return max(0, (ws.max_row or 1) - 1)
        finally:
            wb.close()
    except Exception:
        return 0
