"""Export Excel — fichier en ligne NHR + Statut (consolidé) + Secteur (fichier secteurs)."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from fiscal_constants import ETAT_DEFAILLANT
from models import Contribuable
from services.consolidation import load_consolidated_meta
from services.en_ligne_store import load_en_ligne_headers, load_en_ligne_lines, sectors_applied
from services.stats import get_global_stats

STATUT_COLUMN = 'Statut'
SECTEUR_COLUMN = 'Secteur'
SOUS_SECTEUR_COLUMN = 'Sous secteur'
_EXTRA_COLUMNS = {
    STATUT_COLUMN.lower(), SECTEUR_COLUMN.lower(), SOUS_SECTEUR_COLUMN.lower(),
    'secteur cime', 'sous_secteur', 'sous secteur',
}


def _recherchev_statut_value(contribuable):
    """RECHERCHEV consolidé : montant ou #N/A."""
    if not contribuable or contribuable.etat == ETAT_DEFAILLANT or contribuable.etat is None:
        return '#N/A'
    return float(contribuable.montant_declare or 0)


def _recherchev_secteur_value(contribuable):
    """RECHERCHEV fichier secteurs : libellé Secteur CIME ou #N/A."""
    if not contribuable or not (contribuable.secteur or '').strip():
        return '#N/A'
    return contribuable.secteur.strip()


def _recherchev_sous_secteur_value(contribuable):
    """RECHERCHEV fichier secteurs : sous secteur ou #N/A."""
    if not contribuable or not (contribuable.sous_secteur or '').strip():
        return '#N/A'
    return contribuable.sous_secteur.strip()


def _original_headers(headers):
    """Colonnes d'origine du fichier en ligne (sans Statut / Secteur ajoutés)."""
    return [h for h in headers if h.strip().lower() not in _EXTRA_COLUMNS]


def build_en_ligne_recherchev_export(db, include_secteur=False):
    """
    Fichier en ligne (colonnes d'origine) + Statut [+ Secteur si étape secteurs faite].
    """
    import json

    headers = load_en_ligne_headers(db)
    lines = load_en_ligne_lines(db)
    if not headers or not lines:
        return None, []

    by_niu = {c.niu: c for c in db.query(Contribuable).all()}
    orig = _original_headers(headers)
    export_headers = list(orig)
    export_headers.append(STATUT_COLUMN)
    if include_secteur:
        export_headers.append(SECTEUR_COLUMN)
        export_headers.append(SOUS_SECTEUR_COLUMN)

    rows = []
    for line in lines:
        try:
            values = json.loads(line.values_json or '[]')
        except (json.JSONDecodeError, TypeError):
            values = []
        row_dict = {}
        for i, h in enumerate(headers):
            if h.strip().lower() in _EXTRA_COLUMNS:
                continue
            row_dict[h] = values[i] if i < len(values) else ''
        c = by_niu.get(line.niu)
        row_dict[STATUT_COLUMN] = _recherchev_statut_value(c)
        if include_secteur:
            row_dict[SECTEUR_COLUMN] = _recherchev_secteur_value(c)
            row_dict[SOUS_SECTEUR_COLUMN] = _recherchev_sous_secteur_value(c)
        rows.append(row_dict)

    return export_headers, rows


def _build_synthesis_rows(db, include_secteur=False):
    stats = get_global_stats(db)
    meta = load_consolidated_meta(db) or {}
    rows = [
        {'Indicateur': 'Taille fichier en ligne', 'Valeur': stats['total']},
        {'Indicateur': 'Contribuables dans le consolidé', 'Valeur': meta.get('nb_contribuables', '—')},
        {'Indicateur': 'Fichiers sectoriels consolidés', 'Valeur': meta.get('nb_fichiers', '—')},
        {'Indicateur': 'RECHERCHEV Statut #N/A → défaillants', 'Valeur': stats['defaillants']},
        {'Indicateur': 'Montant 0–10 FCFA → néants', 'Valeur': stats['neants']},
        {'Indicateur': 'Montant > 10 FCFA → déclarants', 'Valeur': stats['declarants']},
    ]
    if include_secteur:
        with_sector = sum(
            1 for c in db.query(Contribuable).all()
            if (c.secteur or '').strip()
        )
        rows.append({'Indicateur': 'RECHERCHEV Secteur trouvés', 'Valeur': with_sector})
        rows.append({'Indicateur': 'RECHERCHEV Secteur #N/A', 'Valeur': stats['total'] - with_sector})
    rows.extend([
        {'Indicateur': 'Taux déclaration (%)', 'Valeur': round(stats['taux_declaration'], 1)},
        {'Indicateur': 'Taux défaillant (%)', 'Valeur': round(stats['taux_defaillant'], 1)},
        {'Indicateur': 'Taux néant (%)', 'Valeur': round(stats['taux_neant'], 1)},
    ])
    return rows


def _write_data_sheet(ws, headers, data_rows):
    header_fill = PatternFill('solid', fgColor='1e40af')
    header_font = Font(bold=True, color='FFFFFF', size=10)
    thin = Side(style='thin', color='cbd5e1')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    na_font = Font(bold=True, color='991b1b')

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border

    statut_idx = headers.index(STATUT_COLUMN) + 1 if STATUT_COLUMN in headers else None
    secteur_idx = headers.index(SECTEUR_COLUMN) + 1 if SECTEUR_COLUMN in headers else None
    sous_idx = headers.index(SOUS_SECTEUR_COLUMN) + 1 if SOUS_SECTEUR_COLUMN in headers else None

    for r_idx, row in enumerate(data_rows, 2):
        for c_idx, h in enumerate(headers, 1):
            val = row.get(h, '')
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.border = border
            if c_idx in (statut_idx, secteur_idx, sous_idx):
                cell.alignment = Alignment(
                    horizontal='right' if c_idx == statut_idx else 'left',
                    vertical='center',
                )
                if val == '#N/A':
                    cell.font = na_font
                elif c_idx == statut_idx and isinstance(val, (int, float)):
                    cell.number_format = '#,##0'
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')

    for i, h in enumerate(headers, 1):
        width = min(max(len(str(h)) + 2, 12), 40)
        hl = h.lower()
        if 'raison' in hl or 'soc' in hl:
            width = 32
        elif h == 'NIU' or 'niu' in hl:
            width = 18
        elif h == SECTEUR_COLUMN or h == SOUS_SECTEUR_COLUMN:
            width = 22
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = 'A2'


def build_vlookup_comparison_workbook(db, include_secteur=None):
    if include_secteur is None:
        include_secteur = sectors_applied(db)
    headers, rows = build_en_ligne_recherchev_export(db, include_secteur=include_secteur)
    if not headers or not rows:
        return None

    wb = Workbook()
    ws = wb.active
    ws.title = 'Fichier en ligne'
    _write_data_sheet(ws, headers, rows)

    ws_synth = wb.create_sheet('Classification')
    _write_data_sheet(ws_synth, ['Indicateur', 'Valeur'], _build_synthesis_rows(db, include_secteur))
    return wb


def vlookup_comparison_to_bytes(db, include_secteur=None):
    wb = build_vlookup_comparison_workbook(db, include_secteur=include_secteur)
    if wb is None:
        return None
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


def export_download_name(include_secteur=False):
    if include_secteur:
        return 'fichier_en_ligne_statut_secteur_sous_secteur.xlsx'
    return 'fichier_en_ligne_statut.xlsx'
