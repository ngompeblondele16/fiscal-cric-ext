"""Export PowerPoint — modèle Evaluation performance CRICEXT.

Remplit les tableaux disponibles depuis les traitements TaxStats.
Laisse vides (espace) les tableaux non encore calculés dans l'app.
"""
from __future__ import annotations

import os
import shutil
import tempfile
from io import BytesIO

try:
    from pptx import Presentation
except ImportError as exc:  # pragma: no cover
    Presentation = None
    _PPTX_IMPORT_ERROR = exc
else:
    _PPTX_IMPORT_ERROR = None

from config import BASE_DIR
from services.ca_analysis import ca_table_rows, get_ca_meta
from services.treatment_wizard import sector_kpi_rows_for_export, sous_secteur_kpi_rows_for_export, reliquataires_sous_secteur_rows_for_export

TEMPLATE_PATH = os.path.join(
    BASE_DIR, 'assets', 'report_templates', 'evaluation_performance_modele.pptx'
)

# Indices 0-based — modèle sans slides 7, 10, 13, 14, 15 (numérotation PowerPoint 1-based)
# Remaining: titre, plan, perf centre, rendement, taux SS, taux secteur,
#            CA activité, reliquataires/émissions SS, chart, recettes struct, merci
SLIDE_PERF_CENTRE = 2          # non disponible → vider
SLIDE_RENDEMENT_FONCTION = 3   # non disponible → vider
SLIDE_TAUX_SOUS_SECTEUR = 4    # disponible
SLIDE_TAUX_SECTEUR = 5         # disponible
SLIDE_CA_ACTIVITE = 6          # disponible (CA)
SLIDE_EMISSIONS = 7            # reliquataires / émissions par sous-secteur
SLIDE_RECETTES_STRUCT = 9      # non disponible → vider

_EMPTY_SLIDES = {
    SLIDE_PERF_CENTRE,
    SLIDE_RENDEMENT_FONCTION,
    SLIDE_RECETTES_STRUCT,
}


def _fmt_int(val):
    try:
        n = int(round(float(val or 0)))
    except (TypeError, ValueError):
        return ''
    return f'{n:,}'.replace(',', ' ')


def _fmt_pct(val, digits=1):
    try:
        return f'{float(val):.{digits}f}%'
    except (TypeError, ValueError):
        return ''


def _fmt_ca(val):
    try:
        if val is None:
            return ''
        return _fmt_int(val)
    except (TypeError, ValueError):
        return ''


def _normalize_key(label):
    return (label or '').strip().upper()


def _find_table(slide):
    for shape in slide.shapes:
        if shape.has_table:
            return shape.table
    return None


def _set_cell_text(cell, text, bold=False):
    """Remplace le texte d'une cellule en conservant un style simple."""
    text = '' if text is None else str(text)
    tf = cell.text_frame
    if not tf.paragraphs:
        cell.text = text
        return
    p = tf.paragraphs[0]
    if p.runs:
        # garder le formatage du premier run
        p.runs[0].text = text
        for run in p.runs[1:]:
            run.text = ''
        if bold:
            p.runs[0].font.bold = True
    else:
        run = p.add_run()
        run.text = text
        if bold:
            run.font.bold = True
    # nettoyer paragraphes suivants
    for extra in tf.paragraphs[1:]:
        for run in extra.runs:
            run.text = ''


def _clear_table_values(table, keep_first_col=True, keep_header=True):
    """Vide les cellules de données (garde en-tête et éventuellement libellés colonne 0)."""
    for ri, row in enumerate(table.rows):
        if keep_header and ri == 0:
            continue
        for ci, cell in enumerate(row.cells):
            if keep_first_col and ci == 0:
                continue
            _set_cell_text(cell, '')


def _row_label_map(table):
    """Map libellé normalisé → index de ligne (hors header)."""
    out = {}
    for ri, row in enumerate(table.rows):
        if ri == 0:
            continue
        label = (row.cells[0].text or '').strip()
        if label:
            out[_normalize_key(label)] = ri
    return out


def _write_row_values(table, row_idx, values_by_col):
    """values_by_col: dict col_index -> text."""
    row = table.rows[row_idx]
    for ci, text in values_by_col.items():
        if ci < len(row.cells):
            _set_cell_text(row.cells[ci], text, bold=(row_idx == len(table.rows) - 1))


def _update_title_slide(prs, period_label):
    if not period_label:
        return
    slide = prs.slides[0]
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        full = shape.text_frame.text or ''
        if 'EVALUATION' in full.upper() and 'PERFORMANCE' in full.upper():
            new_title = f' EVALUATION DE LA PERFORMANCE DU CRICEXT — {period_label} '
            # remplacer le contenu du premier paragraphe
            p = shape.text_frame.paragraphs[0]
            if p.runs:
                p.runs[0].text = new_title
                for run in p.runs[1:]:
                    run.text = ''
            else:
                shape.text_frame.text = new_title
            break


def _fill_sous_secteur_taux(table, rows):
    """Slide 4 — TCD sous-secteur."""
    _clear_table_values(table, keep_first_col=True, keep_header=True)
    label_map = _row_label_map(table)
    total_vals = None

    for r in rows:
        label = (r.get('SOUS SECTEUR') or '').strip()
        key = _normalize_key(label)
        is_total = key.startswith('TOTAL') or key in ('CIMEXT', 'TOTAL GENERAL', 'TOTAL GÉNÉRAL')
        vals = {
            1: _fmt_int(r.get('Taille')),
            2: _fmt_int(r.get('declaration')),
            3: _fmt_pct(r.get('Taux de declaration')),
            4: _fmt_int(r.get('def')) if r.get('def') else '',
            5: _fmt_int(r.get('NEANT')) if r.get('NEANT') else '',
            6: _fmt_int(r.get('Non contributeurs')),
            7: _fmt_pct(r.get('Taux de non contributeurs')),
        }
        if is_total:
            total_vals = vals
            # forcer libellé CIMEXT sur dernière ligne si présente
            last = len(table.rows) - 1
            if last > 0:
                _set_cell_text(table.rows[last].cells[0], 'CIMEXT', bold=True)
                _write_row_values(table, last, vals)
            continue
        ri = label_map.get(key)
        if ri is None and key == '#N/A':
            ri = label_map.get('#N/A') or label_map.get('N/A')
        if ri is not None:
            _write_row_values(table, ri, vals)

    if total_vals is None and rows:
        # fallback: dernière ligne du tableau
        last = len(table.rows) - 1
        if last > 0:
            _set_cell_text(table.rows[last].cells[0], 'CIMEXT', bold=True)


def _fill_secteur_taux(table, rows):
    """Slide 5 — TCD Secteur CIME."""
    _clear_table_values(table, keep_first_col=True, keep_header=True)
    label_map = _row_label_map(table)

    for r in rows:
        label = (r.get('SECTEUR') or '').strip()
        key = _normalize_key(label)
        is_total = key.startswith('TOTAL') or key == 'CIMEXT'
        taille = r.get('Taille du fichier') or 0
        declarant = r.get('DECLARANT') or 0
        defail = r.get('DEFAIL') or 0
        neant = r.get('NEANT') or 0
        non_contrib = int(defail or 0) + int(neant or 0)
        taux_decl = r.get('Taux de declaration')
        taux_nc = round(non_contrib / taille * 100, 1) if taille else 0.0
        vals = {
            1: _fmt_int(taille),
            2: _fmt_int(declarant),
            3: _fmt_pct(taux_decl),
            4: _fmt_int(defail) if defail else '',
            5: _fmt_int(neant) if neant else '',
            6: _fmt_int(non_contrib) if non_contrib else '0',
            7: _fmt_pct(taux_nc),
        }
        if is_total:
            last = len(table.rows) - 1
            _set_cell_text(table.rows[last].cells[0], 'CIMEXT', bold=True)
            _write_row_values(table, last, vals)
            continue
        ri = label_map.get(key)
        if ri is None and key == '#N/A':
            ri = label_map.get('#N/A')
        if ri is not None:
            _write_row_values(table, ri, vals)


def _fill_ca_table(table, ca_rows, info):
    """CA par sous-secteur."""
    _clear_table_values(table, keep_first_col=True, keep_header=True)
    if not ca_rows or not info:
        return

    col_cur = info['col_current']
    col_prev = info['col_previous']
    # Adapter en-têtes si années différentes
    header = table.rows[0]
    if len(header.cells) >= 3:
        _set_cell_text(header.cells[1], col_cur)
        _set_cell_text(header.cells[2], col_prev)

    label_map = _row_label_map(table)
    for r in ca_rows:
        label = (r.get('SOUS SECTEUR') or '').strip()
        key = _normalize_key(label)
        cur = r.get(col_cur) or 0
        prev = r.get(col_prev) or 0
        ecart = r.get('Ecart absolu')
        if ecart is None:
            ecart = float(cur or 0) - float(prev or 0)
        evo = r.get('Evolution')
        evo_txt = _fmt_pct(evo) if evo is not None else ''
        vals = {
            1: _fmt_ca(cur),
            2: _fmt_ca(prev),
            3: _fmt_ca(ecart),
            4: evo_txt,
        }
        if r.get('_is_total') or key in ('TOTAL', 'CIMEXT'):
            last = len(table.rows) - 1
            _set_cell_text(table.rows[last].cells[0], 'CIMEXT', bold=True)
            _write_row_values(table, last, vals)
            continue
        ri = label_map.get(key)
        if ri is None and key == '#N/A':
            ri = label_map.get('#N/A')
        if ri is not None:
            _write_row_values(table, ri, vals)


def _fill_emissions_reliquataires_table(table, rows):
    """Slide émissions — reliquataires et montants émis par sous-secteur."""
    _clear_table_values(table, keep_first_col=True, keep_header=True)
    if not rows:
        return

    header = table.rows[0]
    headers = [
        'SOUS SECTEUR', 'Reliquataires', 'Montant reliquataires',
        'Emissions', 'Ratio des reliquataires',
    ]
    for ci, title in enumerate(headers):
        if ci < len(header.cells):
            _set_cell_text(header.cells[ci], title)

    label_map = _row_label_map(table)
    for r in rows:
        label = (r.get('SOUS SECTEUR') or '').strip()
        key = _normalize_key(label)
        is_total = r.get('_is_total') or key.startswith('TOTAL') or key in ('CIMEXT', 'TOTAL GENERAL', 'TOTAL GÉNÉRAL')
        vals = {
            1: _fmt_int(r.get('Reliquataires')),
            2: _fmt_ca(r.get('Montant reliquataires')),
            3: _fmt_ca(r.get('Emissions')),
            4: _fmt_pct(r.get('Ratio des reliquataires')),
        }
        if is_total:
            last = len(table.rows) - 1
            _set_cell_text(table.rows[last].cells[0], 'CIMEXT', bold=True)
            _write_row_values(table, last, vals)
            continue
        ri = label_map.get(key)
        if ri is None and key == '#N/A':
            ri = label_map.get('#N/A') or label_map.get('N/A')
        if ri is not None:
            _write_row_values(table, ri, vals)


def build_evaluation_pptx(db, period_label=None):
    """
    Génère le PPTX d'évaluation à partir du modèle officiel.
    Retourne (BytesIO, filename) ou lève FileNotFoundError si modèle absent.
    """
    if Presentation is None:
        raise ModuleNotFoundError(
            "Le module python-pptx n'est pas installé. "
            "Exécutez : python -m pip install python-pptx"
        ) from _PPTX_IMPORT_ERROR

    if not os.path.isfile(TEMPLATE_PATH):
        raise FileNotFoundError(
            f'Modèle PowerPoint introuvable : {TEMPLATE_PATH}'
        )

    # Copie temporaire pour ne pas altérer le modèle
    fd, tmp_path = tempfile.mkstemp(suffix='.pptx')
    os.close(fd)
    try:
        shutil.copy2(TEMPLATE_PATH, tmp_path)
        prs = Presentation(tmp_path)

        if period_label is None:
            meta = get_ca_meta(db)
            if meta and meta.year_current:
                period_label = f"Période {meta.year_current}"
        _update_title_slide(prs, period_label)

        # Vider les tableaux non encore produits par l'app
        for idx in _EMPTY_SLIDES:
            if idx < len(prs.slides):
                table = _find_table(prs.slides[idx])
                if table is not None:
                    _clear_table_values(table, keep_first_col=True, keep_header=True)

        # Taux déclaration sous-secteur
        sous_rows = sous_secteur_kpi_rows_for_export(db)
        t4 = _find_table(prs.slides[SLIDE_TAUX_SOUS_SECTEUR])
        if t4 is not None:
            if sous_rows:
                _fill_sous_secteur_taux(t4, sous_rows)
            else:
                _clear_table_values(t4, keep_first_col=True, keep_header=True)

        # Taux déclaration secteur CIME
        sect_rows = sector_kpi_rows_for_export(db)
        t5 = _find_table(prs.slides[SLIDE_TAUX_SECTEUR])
        if t5 is not None:
            if sect_rows:
                _fill_secteur_taux(t5, sect_rows)
            else:
                _clear_table_values(t5, keep_first_col=True, keep_header=True)

        # CA activité
        ca_rows, ca_info = ca_table_rows(db)
        t_ca = _find_table(prs.slides[SLIDE_CA_ACTIVITE])
        if t_ca is not None:
            if ca_rows:
                _fill_ca_table(t_ca, ca_rows, ca_info)
            else:
                _clear_table_values(t_ca, keep_first_col=True, keep_header=True)

        # Reliquataires / émissions par sous-secteur
        rel_rows = reliquataires_sous_secteur_rows_for_export(db)
        t_em = _find_table(prs.slides[SLIDE_EMISSIONS])
        if t_em is not None:
            if rel_rows:
                _fill_emissions_reliquataires_table(t_em, rel_rows)
            else:
                _clear_table_values(t_em, keep_first_col=True, keep_header=True)

        bio = BytesIO()
        prs.save(bio)
        bio.seek(0)

        safe_period = (period_label or 'evaluation').replace('/', '-').replace(' ', '_')
        filename = f'Evaluation_performance_CRICEXT_{safe_period}.pptx'
        return bio, filename
    finally:
        try:
            os.remove(tmp_path)
        except OSError:
            pass
